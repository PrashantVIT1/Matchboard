"""PDF and DOCX text extraction utilities."""

from pathlib import Path
from pypdf import PdfReader
from docx import Document


def read_pdf(file_path: Path) -> str:
    """
    Extract text from PDF file.
    
    Args:
        file_path: Path to PDF file
    
    Returns:
        Extracted text
    """
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_path: Path) -> str:
    """
    Extract text from DOCX file.
    
    Args:
        file_path: Path to DOCX file
    
    Returns:
        Extracted text
    """
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path: Path) -> str | None:
    """
    Extract text from resume file (PDF or DOCX).
    
    Args:
        file_path: Path to resume file
    
    Returns:
        Extracted text or None if unsupported format
    """
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


